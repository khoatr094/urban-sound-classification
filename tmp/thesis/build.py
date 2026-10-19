from pathlib import Path
from copy import deepcopy
import csv,json,zipfile,hashlib,re,datetime
from lxml import etree
from docx import Document
from docx.shared import Cm,Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH,WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from PIL import Image,ImageDraw,ImageFont

ROOT=Path(__file__).resolve().parents[2]
TMP=ROOT/'tmp/thesis'; OUT=ROOT/'output'; OUT.mkdir(exist_ok=True)
REF=Path('D:/OD/KL/Mau bao cao KLTN_DS.docx')
FINAL=OUT/'Maubaocao_KLTN_Phan_loai_am_thanh_do_thi.docx'
source=Document(REF); original=list(source._element.body)
sec=[deepcopy(x) for x in source._element.body.iter(qn('w:sectPr'))]
contract=['# Template contract',f'Reference: {REF}',f'SHA256: {hashlib.sha256(REF.read_bytes()).hexdigest()}',
 'Reference pages: 30 (WPS PDF export inspected in template_contact.jpg). Five source sections.',
 'Retain first 72 body children for three covers; replace names title supervisors month. Preserve source logos drawings borders.',
 'Front matter: source section 2; retain three review forms. Rebuild abstracts gratitude contents lists abbreviations.',
 'Main content: clone Content Heading1 Heading2 Caption patterns. Six chapters explicitly authorized by topic scope.',
 'References: source section 4; appendix source section 5 with no page-number field.',
 'All package parts not edited remain byte-identical through final package merge. Editable document.xml document.xml.rels media settings headers footers.',
 'No signatures grades meeting dates or independent trial results fabricated. Weekly journal is a proposed reconstruction for student confirmation.',
 'Renderer: packaged render_docx.py failed because LibreOffice is absent on Windows. Native WPS COM export used for page PNG review.']
for i,s in enumerate(source.sections):
 contract.append(f'Section {i+1}: page {s.page_width/360000:.4f}x{s.page_height/360000:.4f} cm; margins T{s.top_margin/360000:.4f} B{s.bottom_margin/360000:.4f} L{s.left_margin/360000:.4f} R{s.right_margin/360000:.4f}; header {s.header_distance/360000:.4f} footer {s.footer_distance/360000:.4f}.')
for name in ['Normal','Content','Heading 1','Heading 2','Caption']:
 st=source.styles[name]; contract.append(name+': '+etree.tostring(st.element,encoding='unicode'))
(TMP/'artifact.md').write_text('\n'.join(contract),encoding='utf8')
inventory={n:{'bytes':len(b),'sha256':hashlib.sha256(b).hexdigest()} for n in zipfile.ZipFile(REF).namelist() for b in [zipfile.ZipFile(REF).read(n)]}
(TMP/'package_inventory.json').write_text(json.dumps(inventory,indent=2),encoding='utf8')

def font(sz=34,bold=False):
 return ImageFont.truetype('C:/Windows/Fonts/'+('timesbd.ttf' if bold else 'times.ttf'),sz)
def diagram(name,rows):
 im=Image.new('RGB',(1800,len(rows)*165+65),'white'); dr=ImageDraw.Draw(im)
 for i,(title,detail) in enumerate(rows):
  top=30+i*165; dr.rounded_rectangle((80,top,1720,top+115),radius=12,outline='#333333',width=3,fill='#f8f8f8')
  dr.text((900,top+27),title,font=font(37,True),anchor='mm',fill='black')
  dr.text((900,top+77),detail,font=font(32),anchor='mm',fill='black')
  if i<len(rows)-1:
   dr.line((900,top+116,900,top+151),fill='black',width=4)
   dr.polygon([(890,top+142),(910,top+142),(900,top+155)],fill='black')
 im.save(TMP/name)
diagram('pipeline.png',[
 ('Âm thanh và metadata','Kiểm tra giải mã · SHA-256 · fsID · lọc 7 lớp'),
 ('Chia theo fold và loại nguồn giao nhau','Train 4.632 · Validation 553 · Test 541'),
 ('Hai nhánh đặc trưng','CNN: Log-Mel 64 + chuẩn hóa train | AST: bộ xử lý pretrained'),
 ('Huấn luyện và chọn trên validation','CNN4 · ResCNN · AST | chọn checkpoint và trọng số ensemble'),
 ('Đánh giá và ứng dụng','Test fold 10 → CSV, hình vẽ | Tệp người dùng → xác suất 7 lớp')])
diagram('cnn4.png',[
 ('Đầu vào','Log-Mel 64 dải · 1 kênh · chuẩn hóa theo dải Mel'),
 ('Bốn khối CNN','Kênh 32 → 64 → 128 → 256'),
 ('Cấu trúc mỗi khối','2 × Conv 3×3 + BN + ReLU → MaxPool 2×2 → Dropout2d'),
 ('Gom đặc trưng và phân loại','Global average pooling → Linear 256→128 → ReLU → Dropout 0,35'),
 ('Đầu ra','Linear 128→7 → softmax khi suy luận')])
diagram('rescnn.png',[
 ('Stem','Conv 3×3, 1→32 → BatchNorm → GELU'),
 ('Bốn khối phần dư','32→64 → 128 → 256 → 256 · stride 2 mỗi khối'),
 ('Nhánh chính và nhánh tắt','Conv-BN-GELU → Conv-BN → SE → Dropout + phép chiếu skip'),
 ('Tổng hợp đặc trưng','Global mean 256 chiều + global max 256 chiều → ghép 512'),
 ('Phân loại','LayerNorm → Dropout → Linear 512→128 → GELU → Linear 128→7')])
diagram('ast.png',[
 ('Tín hiệu mono 16 kHz','Cửa sổ 4 giây · ASTFeatureExtractor của checkpoint MIT'),
 ('Đặc trưng AST','Tối đa 1.024 khung × 128 dải Mel · chuẩn hóa pretrained'),
 ('Biểu diễn patch','Chiếu patch thành vector + thông tin vị trí + token đặc biệt'),
 ('Các tầng Transformer','Multi-head self-attention → mạng truyền thẳng → biểu diễn toàn cục'),
 ('Tinh chỉnh có giám sát','Thay đầu phân loại 7 lớp · AdamW · batch 2 × tích lũy 8 bước')])

d=Document(REF)
for el in list(d._element.body): d._element.body.remove(el)
for el in original[:72]: d._element.body.append(deepcopy(el))
# Retain terminal geometry while python-docx adds the replacement content.
d._element.body.append(deepcopy(sec[-1]))
def replace(p,text):
 rpr=deepcopy(p.runs[0]._r.rPr) if p.runs and p.runs[0]._r.rPr is not None else None
 for el in list(p._p):
  if el.tag!=qn('w:pPr'): p._p.remove(el)
 r=p.add_run(text)
 if rpr is not None:r._r.insert(0,rpr)
for p in d.paragraphs:
 t=p.text
 if 'HỌ VÀ TÊN SINH VIÊN 1' in t: replace(p,'LƯƠNG CÔNG PHÚ BOY - 22002515')
 elif 'HỌ VÀ TÊN SINH VIÊN 2' in t: replace(p,'TRẦN ANH KHOA - 22000915')
 elif 'HƯỚNG DẪN VÀ MẪU BÁO CÁO' in t:replace(p,'PHÂN LOẠI ÂM THANH ĐÔ THỊ\nSỬ DỤNG DEEP LEARNING')
 elif 'TS. Họ và Tên GVHD' in t:replace(p,'Giảng viên hướng dẫn: TS. Huỳnh Công Bằng\nTS. Nguyễn Minh Hạnh')
 elif 'THÁNG 12 NĂM 2026' in t:replace(p,t.replace('THÁNG 12','THÁNG 10'))
 elif 'FULLNAME OF THE FIRST' in t:replace(p,'LUONG CONG PHU BOY - 22002515')
 elif 'FULLNAME OF THE SECOND' in t:replace(p,'TRAN ANH KHOA - 22000915')
 elif 'GUIDANCE AND SAMPLE' in t:replace(p,'URBAN SOUND CLASSIFICATION\nUSING DEEP LEARNING')
 elif 'Supervisor:' in t:replace(p,'Supervisors: Dr. Huynh Cong Bang\nDr. Nguyen Minh Hanh')
 elif 'DECEMBER' in t:replace(p,t.replace('DECEMBER','OCTOBER'))

def p(text='',style='Content',align=None):
 x=d.add_paragraph(text,style)
 if align is not None:x.alignment=align
 x.paragraph_format.widow_control=True
 return x
def front(title,new=True):
 if new:d.add_page_break()
 x=p(title,'Heading 1N',WD_ALIGN_PARAGRAPH.CENTER)
 x.paragraph_format.space_before=Pt(12);x.paragraph_format.space_after=Pt(18)
 for r in x.runs:r.bold=True;r.font.size=Pt(14)
 return x
def endsection(s):
 x=p('','Normal'); x.paragraph_format.space_before=Pt(0);x.paragraph_format.space_after=Pt(0)
 x._p.get_or_add_pPr().append(deepcopy(s))
def bookmark(x,name):
 b=OxmlElement('w:bookmarkStart');b.set(qn('w:id'),str(bookmark.counter));b.set(qn('w:name'),name)
 e=OxmlElement('w:bookmarkEnd');e.set(qn('w:id'),str(bookmark.counter));bookmark.counter+=1
 x._p.insert(1,b);x._p.append(e)
bookmark.counter=1000
headings=[]; figs=[]; tables=[]; toc_slots=[]; figure_slots=[];table_slots=[]
pagecache=json.loads((TMP/'pagecache.json').read_text(encoding='utf8')) if (TMP/'pagecache.json').exists() else {}
def listline(text,key,level=0):
 x=p(text+'\t'+str(pagecache.get(key,'…')),'Normal')
 x.paragraph_format.line_spacing=1.15;x.paragraph_format.space_before=Pt(0);x.paragraph_format.space_after=Pt(3)
 x.paragraph_format.left_indent=Cm(.45*level)
 from docx.enum.text import WD_TAB_ALIGNMENT,WD_TAB_LEADER
 x.paragraph_format.tab_stops.add_tab_stop(Cm(15.0-.45*level),WD_TAB_ALIGNMENT.RIGHT,WD_TAB_LEADER.DOTS)
 return x

front('URBAN SOUND CLASSIFICATION USING DEEP LEARNING',False)
p('ABSTRACT','Normal',WD_ALIGN_PARAGRAPH.CENTER).runs[0].bold=True
p('This thesis investigates deep learning methods for classifying seven urban sound categories selected from UrbanSound8K: car horn, children playing, dog bark, engine idling, gunshot, siren and drilling. The study compares a four-block convolutional network, a residual CNN with channel attention, and an Audio Spectrogram Transformer initialized from an AudioSet-pretrained checkpoint. A weighted probability ensemble is also evaluated.')
p('The experimental pipeline checks audio validity and duplicate recordings, converts signals to mono at 16 kHz, and separates training, validation and test data by the predefined folds. After removing six training excerpts sharing recording sources with higher-priority partitions, the sets contain 4,632, 553 and 541 excerpts, respectively. CNN features are normalized using training-only statistics for each Mel band. Checkpoints and ensemble weights are selected on validation macro-F1.')
p('On test fold 10, AST achieves 92.05% accuracy and 91.88% macro-F1. CNN4 reaches 85.58% accuracy and 87.23% macro-F1, while ResCNN reaches 83.73% and 86.11%. The ensemble matches AST accuracy but does not improve test macro-F1. Error analysis identifies siren-to-dog-bark and drilling-to-car-horn confusions as the dominant error groups. A local application provides audio upload, recording and probability visualization.')
p('These findings are limited to a seven-class, single-fold experiment. They are not ten-fold UrbanSound8K benchmark results. Broader evaluation requires multiple folds and seeds, an independent out-of-class dataset, and a version-locked deployment package.')
p('Keywords: urban sound classification, deep learning, Log-Mel spectrogram, convolutional neural network, Audio Spectrogram Transformer.','Content')
front('TÓM TẮT')
p('Khóa luận nghiên cứu bài toán phân loại bảy nhóm âm thanh đô thị được chọn từ UrbanSound8K, gồm còi xe, trẻ em vui chơi, chó sủa, động cơ chạy không tải, tiếng súng, còi báo động và tiếng khoan. Nhóm so sánh CNN4, ResCNN có chú ý kênh và Audio Spectrogram Transformer được khởi tạo từ checkpoint pretrained trên AudioSet. Phương án ensemble xác suất có trọng số được đánh giá như một cách kết hợp các mô hình.')
p('Quy trình thực nghiệm kiểm tra tệp âm thanh, mã băm và nguồn ghi âm trước khi tạo đặc trưng. Tín hiệu được đưa về mono 16 kHz với cửa sổ 4 giây. Sau khi loại sáu mẫu train giao nguồn với tập ưu tiên cao hơn, train có 4.632 mẫu, validation 553 mẫu và test 541 mẫu. Thống kê chuẩn hóa CNN được tính riêng theo dải Mel trên train. Checkpoint và trọng số kết hợp được chọn bằng validation macro-F1.')
p('Trên test fold 10, AST đạt accuracy 92,05% và macro-F1 91,88%. CNN4 đạt 85,58% và 87,23%; ResCNN đạt 83,73% và 86,11%. Ensemble có accuracy bằng AST nhưng chưa cải thiện macro-F1 trên test. Phân tích lỗi cho thấy hai nhóm nhầm lẫn đáng chú ý là còi báo động thành chó sủa và tiếng khoan thành còi xe. Ứng dụng cục bộ hỗ trợ tải tệp, ghi âm và xem xác suất theo lớp.')
p('Kết quả phản ánh một phép chia bảy lớp, chưa phải đánh giá chéo mười fold của benchmark gốc. Hướng phát triển tập trung vào đánh giá nhiều fold, nhiều seed, dữ liệu ngoài lớp và thống nhất gói triển khai với phiên bản đã huấn luyện.')
p('Từ khóa: phân loại âm thanh đô thị, học sâu, phổ Log-Mel, mạng tích chập, Audio Spectrogram Transformer.')
front('LỜI CẢM ƠN')
p('Nhóm chúng em xin gửi lời cảm ơn đến TS. Huỳnh Công Bằng và TS. Nguyễn Minh Hạnh đã hướng dẫn đề tài khóa luận tốt nghiệp “Phân loại âm thanh đô thị sử dụng Deep Learning”. Sự hướng dẫn của thầy, cô là cơ sở để nhóm định hướng nội dung nghiên cứu và hoàn thiện báo cáo.')
p('Chúng em cảm ơn các thầy, cô Khoa Công nghệ Thông tin, Trường Đại học Công nghiệp Thành phố Hồ Chí Minh đã giảng dạy những kiến thức nền tảng về lập trình, xử lý dữ liệu và học máy. Những kiến thức này giúp nhóm tiếp cận tài liệu chuyên ngành, tổ chức thực nghiệm và phân tích kết quả của đề tài.')
p('Nhóm cũng trân trọng công sức của các tác giả UrbanSound8K và những công trình, thư viện nguồn mở được sử dụng. Việc công bố dữ liệu, mã nguồn và mô hình đã tạo điều kiện để sinh viên có thể tìm hiểu và kiểm tra các phương pháp học sâu trên bài toán âm thanh.')
p('Trong phạm vi thời gian và kinh nghiệm còn hạn chế, báo cáo khó tránh khỏi thiếu sót. Nhóm mong nhận được ý kiến của thầy, cô để tiếp tục hoàn thiện phương pháp và cách trình bày kết quả.')
p('Nhóm sinh viên thực hiện','Content',WD_ALIGN_PARAGRAPH.RIGHT)
p('Lương Công Phú Boy\nTrần Anh Khoa','Content',WD_ALIGN_PARAGRAPH.RIGHT)
# Preserve review-page structures and blank signature dates from the supplied template.
d.add_page_break()
for el in original[116:185]: d._element.body.insert(len(d._element.body)-1,deepcopy(el))

lines=(TMP/'content.md').read_text(encoding='utf8').splitlines()
chapter=0;sub=0
outline=[]
for line in lines:
 if line.startswith('# '):
  chapter+=1;sub=0;outline.append((f'CHƯƠNG {chapter}  '+line[2:],f'ch{chapter}',0))
 elif line.startswith('## '):
  sub+=1;outline.append((f'{chapter}.{sub}  '+line[3:],f'h{chapter}_{sub}',1))
 elif line.startswith('!FIG|'):
  _,num,path,title=line.split('|');figs.append((f'Hình {num} {title}',f'fig{num.replace(".","_")}'))
 elif line.startswith('!TABLE|'):
  _,num,title,kind=line.split('|');tables.append((f'Bảng {num} {title}',f'tab{num.replace(".","_")}'))

front('MỤC LỤC')
for title,key,level in outline:listline(title,key,level)
for title,key in [('TÀI LIỆU THAM KHẢO','refs'),('PHỤ LỤC','appendix'),('NHẬT KÝ LÀM VIỆC','journal')]:listline(title,key)
front('DANH MỤC HÌNH ẢNH')
for title,key in figs:listline(title,key)
front('DANH MỤC BẢNG BIỂU')
for title,key in tables:listline(title,key)
listline('Bảng A.1 Các tệp dùng để kiểm tra số liệu','tabA1')
listline('Bảng B.1 Nhật ký công việc theo tuần','tabB1')

def table(headers,rows,widths=None,size=11):
 t=d.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 if widths:
  for col,w in zip(t.columns,widths):col.width=Cm(w)
 for i,h in enumerate(headers):t.rows[0].cells[i].text=h
 for row in rows:
  for c,txt in zip(t.add_row().cells,row):c.text=str(txt)
 for ri,row in enumerate(t.rows):
  pr=row._tr.get_or_add_trPr();cs=OxmlElement('w:cantSplit');pr.append(cs)
  if ri==0:
   rep=OxmlElement('w:tblHeader');pr.append(rep)
  for ci,cell in enumerate(row.cells):
   if widths:cell.width=Cm(widths[ci])
   cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   tcpr=cell._tc.get_or_add_tcPr();mar=OxmlElement('w:tcMar')
   for side in ['top','left','bottom','right']:
    n=OxmlElement('w:'+side);n.set(qn('w:w'),'70');n.set(qn('w:type'),'dxa');mar.append(n)
   tcpr.append(mar)
   if ri==0:
    shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'EEEEEE');tcpr.append(shade)
   for x in cell.paragraphs:
    x.style=d.styles['Normal'];x.paragraph_format.line_spacing=1.05
    x.paragraph_format.left_indent=Cm(0);x.paragraph_format.right_indent=Cm(0);x.paragraph_format.first_line_indent=Cm(0)
    x.paragraph_format.space_before=Pt(2);x.paragraph_format.space_after=Pt(2)
    if len(rows)<16:x.paragraph_format.keep_with_next=ri<len(t.rows)-1
    x.alignment=WD_ALIGN_PARAGRAPH.LEFT if ci==0 else WD_ALIGN_PARAGRAPH.CENTER
    for r in x.runs:r.font.name='Times New Roman';r.font.size=Pt(size);r.bold=ri==0
 borders=OxmlElement('w:tblBorders')
 for side in ['top','left','bottom','right','insideH','insideV']:
  n=OxmlElement('w:'+side);n.set(qn('w:val'),'single');n.set(qn('w:sz'),'4');n.set(qn('w:color'),'777777');borders.append(n)
 t._tbl.tblPr.append(borders)
 return t
front('DANH MỤC CÁC THUẬT NGỮ VIẾT TẮT')
table(['Viết tắt','Thuật ngữ','Nghĩa'],[
 ('AMP','Automatic Mixed Precision','Độ chính xác hỗn hợp tự động'),('AST','Audio Spectrogram Transformer','Transformer trên phổ âm thanh'),
 ('AUC','Area Under the Curve','Diện tích dưới đường cong'),('BN','Batch Normalization','Chuẩn hóa theo batch'),
 ('CNN','Convolutional Neural Network','Mạng nơ-ron tích chập'),('CPU','Central Processing Unit','Bộ xử lý trung tâm'),
 ('CSV','Comma-Separated Values','Dữ liệu phân tách bằng dấu phẩy'),('GPU','Graphics Processing Unit','Bộ xử lý đồ họa'),
 ('PCA','Principal Component Analysis','Phân tích thành phần chính'),('ReLU','Rectified Linear Unit','Hàm kích hoạt tuyến tính chỉnh lưu'),
 ('ROC','Receiver Operating Characteristic','Đường đặc trưng hoạt động'),('SE','Squeeze-and-Excitation','Cơ chế chú ý theo kênh'),
 ('STFT','Short-Time Fourier Transform','Biến đổi Fourier ngắn hạn')],[1.8,6.7,6.5])
endsection(sec[1])

def caption(text,key):
 x=p(text,'Caption',WD_ALIGN_PARAGRAPH.CENTER);bookmark(x,key)
 x.paragraph_format.space_before=Pt(6);x.paragraph_format.space_after=Pt(6)
 x.paragraph_format.first_line_indent=Cm(0)
 x.paragraph_format.keep_together=True
 x.paragraph_format.keep_with_next=True
 for r in x.runs:r.font.size=Pt(13);r.italic=False
 return x
def figure(num,path,title):
 ip=ROOT/path;im=Image.open(ip);w,h=im.size
 maxw=15.5;maxh=13.7
 if path.endswith('05_learning_curves_comparison.png'):maxh=9.5
 if path.endswith('waveform_clean.png'):maxh=20.0
 scale=min(maxw/w,maxh/h)
 x=p('','Normal',WD_ALIGN_PARAGRAPH.CENTER);x.paragraph_format.keep_with_next=True
 x.paragraph_format.first_line_indent=Cm(0);x.paragraph_format.left_indent=Cm(0);x.paragraph_format.right_indent=Cm(0)
 x.paragraph_format.line_spacing=1.0
 x.paragraph_format.space_before=Pt(8);x.paragraph_format.space_after=Pt(0)
 x.add_run().add_picture(str(ip),width=Cm(w*scale),height=Cm(h*scale))
 caption(f'Hình {num} {title}',f'fig{num.replace(".","_")}')
 # Single compact provenance statement keeps captions readable and traceable.
 note='Nguồn: nhóm xây dựng theo notebook thực nghiệm.' if path.startswith('tmp/') and not path.endswith('_clean.png') else 'Nguồn: kết quả thực nghiệm của nhóm.'
 x=p(note,'Normal',WD_ALIGN_PARAGRAPH.CENTER)
 x.paragraph_format.first_line_indent=Cm(0)
 x.paragraph_format.line_spacing=1.0;x.paragraph_format.space_after=Pt(6)
 for r in x.runs:r.font.size=Pt(10);r.italic=True

def readcsv(path):return list(csv.DictReader((ROOT/path).open(encoding='utf-8-sig')))
labels=json.loads((ROOT/'configs/labels.json').read_text(encoding='utf8'))
def fmt(v,n=2):return f'{float(v):.{n}f}'.replace('.',',')
def result_table(num,title,kind):
 x=caption(f'Bảng {num} {title}',f'tab{num.replace(".","_")}');x.paragraph_format.keep_with_next=True
 if kind=='labels':
  table(['Mã','Nhãn gốc','Tên tiếng Việt'],[(i,en,vi) for i,(en,vi) in enumerate(zip(labels['classes'],labels['labels_vi']))],[1.3,6.1,7.6],12)
 elif kind=='split':
  rows=readcsv('outputs/data/split_class_distribution.csv');table(['Lớp','Train','Validation','Test'],[list(r.values()) for r in rows]+[['Tổng','4.632','553','541']],[7.0,2.6,2.8,2.6],12)
 elif kind=='training':
  table(['Thông số','CNN4 và ResCNN','AST'],[
   ('Đầu vào','Log-Mel 64 dải','ASTFeatureExtractor 128 dải'),('Tần số và cửa sổ','16 kHz; 4 giây','16 kHz; 4 giây'),
   ('Batch vật lý','16','2'),('Tích lũy gradient','1 bước','8 bước'),('Epoch tối đa','45','12'),
   ('Optimizer','AdamW','AdamW'),('Tốc độ học ban đầu','0,0008','0,00001'),('Weight decay','0,0005','0,0005'),
   ('Lịch tốc độ học','Warm-up 10% + cosine','Warm-up 10% + cosine'),('Label smoothing','0,06','0,06'),
   ('Patience','10','10'),('Seed','42','42'),('Chọn checkpoint','Validation macro-F1','Validation macro-F1')],[5.3,4.85,4.85],11)
 elif kind=='results':
  rows=readcsv('outputs/metrics/test_summary_metrics.csv')
  table(['Mô hình','Val F1','Test Acc','Macro P','Macro R','Macro F1','W. F1'],[[r['Mã'].upper()]+[fmt(v) for k,v in r.items() if k not in ['Mã','Mô hình']] for r in rows],[2.6,2.05,2.05,2.05,2.05,2.05,2.15],10)
  p('Các chỉ số trong bảng tính theo %. P: precision; R: recall; W. F1: weighted-F1.','Normal')
 elif kind=='perclass':
  rows=[r for r in readcsv('outputs/metrics/per_class_detailed_metrics.csv') if r['Mô hình']=='AST Pretrained']
  counts=[33,100,100,93,32,83,100]
  table(['Lớp','Số mẫu','Precision (%)','Recall (%)','F1 (%)'],[[r['Lớp âm thanh'],n]+[fmt(list(r.values())[i]) for i in [2,3,4]] for r,n in zip(rows,counts)],[5.6,1.6,2.6,2.6,2.6],11)
 elif kind=='cost':
  rows=readcsv('outputs/metrics/benchmark_fps_stability_accuracy.csv')
  table(['Mô hình','Tham số (triệu)','Checkpoint (MiB)','Train (s)','Clip/giây'],[[r['Mô hình'],fmt(r['Tham số (Triệu)'],3),fmt(r['Dung lượng (MB)'],3),fmt(r['Thời gian huấn luyện (s)']),fmt(r['Tốc độ suy luận (FPS)'])] for r in rows],[4.0,2.5,2.9,2.8,2.8],11)

chapter=0;sub=0
for line in lines:
 if not line.strip():continue
 if line.startswith('# '):
  chapter+=1;sub=0
  x=p(f'CHƯƠNG {chapter}\n'+line[2:],'Heading 1');x.paragraph_format.page_break_before=chapter>1
  x.paragraph_format.keep_with_next=True;bookmark(x,f'ch{chapter}')
  np=OxmlElement('w:numPr');ni=OxmlElement('w:numId');ni.set(qn('w:val'),'0');np.append(ni);x._p.get_or_add_pPr().append(np)
 elif line.startswith('## '):
  sub+=1;x=p(f'{chapter}.{sub}  '+line[3:],'Heading 2');bookmark(x,f'h{chapter}_{sub}')
  np=OxmlElement('w:numPr');ni=OxmlElement('w:numId');ni.set(qn('w:val'),'0');np.append(ni);x._p.get_or_add_pPr().append(np)
  x.paragraph_format.keep_with_next=True
 elif line.startswith('!FIG|'):
  _,num,path,title=line.split('|');figure(num,path,title)
 elif line.startswith('!TABLE|'):
  _,num,title,kind=line.split('|');result_table(num,title,kind)
 elif line.startswith('!EQ|'):
  x=p(line[4:],'Normal',WD_ALIGN_PARAGRAPH.CENTER)
  x.paragraph_format.space_before=Pt(10);x.paragraph_format.space_after=Pt(8)
 else:p(line)

endsection(sec[2])
front('TÀI LIỆU THAM KHẢO',False);bookmark(d.paragraphs[-1],'refs')
p('Tiếng Anh','Normal').runs[0].bold=True
refs=[
 '[1] Gong, Y., Chung, Y.-A., Glass, J. (2021), “AST: Audio Spectrogram Transformer”, Proceedings of Interspeech 2021, tr. 571–575. DOI: 10.21437/Interspeech.2021-698. https://arxiv.org/abs/2104.01778',
 '[2] He, K., Zhang, X., Ren, S., Sun, J. (2015), “Deep Residual Learning for Image Recognition”, arXiv:1512.03385. https://arxiv.org/abs/1512.03385',
 '[3] Hu, J., Shen, L., Albanie, S., Sun, G., Wu, E. (2019), “Squeeze-and-Excitation Networks”, arXiv:1709.01507v4, bản mở rộng của công trình CVPR 2018. https://arxiv.org/abs/1709.01507',
 '[4] MIT (không ghi năm), “ast-finetuned-audioset-10-10-0.4593”, thẻ mô hình và checkpoint trên Hugging Face. https://huggingface.co/MIT/ast-finetuned-audioset-10-10-0.4593 (truy cập 19/10/2026).',
 '[5] Park, D. S., Chan, W., Zhang, Y., Chiu, C.-C., Zoph, B., Cubuk, E. D., Le, Q. V. (2019), “SpecAugment: A Simple Data Augmentation Method for Automatic Speech Recognition”, Proceedings of Interspeech 2019, tr. 2613–2617. DOI: 10.21437/Interspeech.2019-2680. https://arxiv.org/abs/1904.08779',
 '[6] Salamon, J., Jacoby, C., Bello, J. P. (2014), “A Dataset and Taxonomy for Urban Sound Research”, Proceedings of the 22nd ACM International Conference on Multimedia, tr. 1041–1044. DOI: 10.1145/2647868.2655045. https://www.justinsalamon.com/uploads/4/3/9/4/4394963/salamon_urbansound_acmmm14.pdf',
 '[7] Salamon, J., Bello, J. P. (2016), “Deep Convolutional Neural Networks and Data Augmentation for Environmental Sound Classification”, arXiv:1608.04363. https://arxiv.org/abs/1608.04363',
 '[8] Urban Sound Datasets (không ghi năm), “UrbanSound8K”, mô tả dữ liệu, metadata và khuyến nghị đánh giá. https://urbansounddataset.weebly.com/urbansound8k.html (truy cập 19/10/2026).',
 '[9] Walden, F., Dasgupta, S., Rahman, M., Islam, M. (2022), “Improving the Environmental Perception of Autonomous Vehicles using Deep Learning-based Audio Classification”, arXiv:2209.04075. https://arxiv.org/abs/2209.04075'
]
for text in refs:
 x=p(text);x.paragraph_format.left_indent=Cm(1);x.paragraph_format.first_line_indent=Cm(-1)
 x.paragraph_format.keep_together=True
endsection(sec[3])
front('PHỤ LỤC',False);bookmark(d.paragraphs[-1],'appendix')
p('A  DỮ LIỆU VÀ KHẢ NĂNG TÁI LẬP','Heading 2')
p('Các đường dẫn dưới đây tương đối với thư mục dự án urban_sound. Bảng tập hợp tệp cần thiết để kiểm tra thông số và kết quả. Giá trị trong báo cáo được làm tròn khi trình bày; tệp CSV giữ giá trị số gốc.')
caption('Bảng A.1 Các tệp dùng để kiểm tra số liệu','tabA1').paragraph_format.keep_with_next=True
table(['Tệp hoặc thư mục','Nội dung đối chiếu'],[
 ('urbansoundclassification.ipynb','Mã đã chạy; môi trường; cấu hình; đầu ra từng cell'),
 ('outputs/data/audit.json','Số mẫu hợp lệ, tệp lỗi, nguồn giao fold và trùng băm'),
 ('outputs/data/split_fold10.json','Fold và số mẫu sau loại nguồn trùng'),
 ('outputs/data/split_class_distribution.csv','Số mẫu từng lớp trên ba tập'),
 ('outputs/data/cnn_norm_fold10.npz','Thống kê chuẩn hóa theo dải Mel'),
 ('outputs/logs/*_history.csv','Loss, accuracy, F1 và thời gian qua epoch'),
 ('outputs/metrics/test_summary_metrics.csv','Chỉ số tổng thể của bốn phương án'),
 ('outputs/metrics/per_class_detailed_metrics.csv','Precision, recall và F1 từng lớp'),
 ('outputs/metrics/benchmark_fps_stability_accuracy.csv','Tham số, dung lượng, thời gian và throughput'),
 ('outputs/checkpoints/*_best.pt','Trọng số đã lưu của từng mô hình'),
 ('configs/ast_processor/; configs/labels.json','Bộ xử lý AST và thứ tự bảy nhãn')],[8,7],10)
p('A.1  Quy trình kiểm tra lại kết quả','Heading 2')
p('Đầu tiên, đối chiếu ba tổng số mẫu với split_fold10.json và bảng phân bố lớp. Tiếp theo, kiểm tra thứ tự nhãn trong configs/labels.json trước khi đọc các ma trận. Sau đó, dùng test_summary_metrics.csv để dựng bảng tổng thể và per_class_detailed_metrics.csv để dựng bảng theo lớp. Khi vẽ lại learning curves, lấy trực tiếp history của mô hình tương ứng, không suy ra từ hình PNG đã nén.')
p('Lệnh chạy ứng dụng tại thư mục dự án: .\\run_windows.bat. Điểm vào Python là app.py. Các lệnh dự đoán và vẽ biểu đồ được lưu tại scripts/predict.py và scripts/plot_results.py. Cần đối chiếu đường dẫn đầu ra mà script hỗ trợ với cấu trúc checkpoint thực tế trước khi chạy lại. Không dùng việc chương trình mở được giao diện làm bằng chứng rằng số liệu thực nghiệm đã được tái lập.')
p('A.2  Thông số cần giữ nhất quán','Heading 2')
p('Bản notebook được phân tích dùng CustomCNN4 hai tích chập 3 × 3 mỗi khối và CustomResCNN có SE. Các tệp cnn4_architecture.txt và rescnn_architecture.txt trong configs mô tả một biến thể cũ hơn. Tương tự, thống kê toàn cục trong stats_fold10.json không thay thế thống kê theo dải ở cnn_norm_fold10.npz. Khi bàn giao cần đóng gói cấu hình khớp tên và kích thước tensor của checkpoint, kèm checksum để phân biệt phiên bản.')
p('Bộ trọng số ensemble tối ưu và xác suất validation từng mẫu cần được xuất cùng bản notebook để tái lập phép kết hợp. Nếu ứng dụng không tìm thấy ensemble.json và dùng trung bình đều, kết quả được hiểu là ensemble demo. Bảng 5.1 không được gán cho cấu hình trung bình đều nếu chưa kiểm tra chúng trùng nhau.')
p('A.3  Kịch bản kiểm tra trước buổi trình bày','Heading 2')
p('Chuẩn bị ít nhất một tệp WAV rõ nguồn, một tệp im lặng và một tệp không hợp lệ. Với tệp hợp lệ, kiểm tra tổng bảy xác suất gần bằng 1, thứ tự tên lớp, thời lượng hiển thị và mô hình đã chọn. Với tệp im lặng hoặc lỗi, kiểm tra thông báo phù hợp và giao diện còn sử dụng được. Với tệp dài, đo riêng thời gian tiền xử lý, thời gian mô hình và tổng thời gian phía ứng dụng.')
p('Phần trình bày slide nên đi theo mạch: bài toán, phạm vi bảy lớp, dữ liệu và kiểm soát nguồn, ba kiến trúc, cấu hình huấn luyện, kết quả tổng thể, lỗi theo lớp, demo và hạn chế. Các biểu đồ dùng trên slide cần giữ nhãn test fold 10. Không đổi nhãn thành kết quả mười fold khi chưa thực hiện đủ các lượt đánh giá.')
front('NHẬT KÝ LÀM VIỆC');bookmark(d.paragraphs[-1],'journal')
p('Đề tài: Phân loại âm thanh đô thị sử dụng Deep Learning')
p('Sinh viên: Lương Công Phú Boy - 22002515; Trần Anh Khoa - 22000915.')
p('Giảng viên hướng dẫn: TS. Huỳnh Công Bằng; TS. Nguyễn Minh Hạnh.')
p('Thời gian tổng hợp: 16/6/2026 đến 19/10/2026.')
p('Bảng tiến độ dự thảo để hai sinh viên đối chiếu và xác nhận theo lịch thực hiện. Ngày báo cáo và nhận xét của giảng viên được để trống.','Normal')
caption('Bảng B.1 Nhật ký công việc theo tuần','tabB1').paragraph_format.keep_with_next=True
tasks=[
 'Khảo sát bài toán phân loại âm thanh đô thị.\nLập đề cương và phạm vi đề tài.',
 'Tìm hiểu UrbanSound8K và nguồn dữ liệu.\nĐọc tài liệu CNN cho âm thanh môi trường.',
 'Khảo sát metadata, nhãn và cấu trúc fold.\nTổng hợp tài liệu tham khảo ban đầu.',
 'Xây dựng mã đọc và kiểm tra tệp âm thanh.\nViết phần giới thiệu bài toán.',
 'Thực hiện chuyển mono, resample và cửa sổ 4 giây.\nKhảo sát waveform và Log-Mel.',
 'Thiết kế chia train, validation, test theo fold.\nKiểm tra fsID và mã băm tệp.',
 'Xây dựng CNN4 và pipeline đặc trưng.\nSoạn phần cơ sở lý thuyết CNN.',
 'Xây dựng ResCNN với kết nối phần dư và SE.\nBổ sung sơ đồ kiến trúc vào báo cáo.',
 'Tích hợp bộ xử lý và checkpoint AST.\nTìm hiểu cơ chế fine-tuning và AMP.',
 'Hoàn thiện tăng cường dữ liệu và hàm mất mát.\nKiểm tra quy trình chọn checkpoint theo validation.',
 'Tổ chức nhật ký loss, accuracy và macro-F1.\nHoàn thiện mã lưu kết quả thực nghiệm.',
 'Xây dựng ensemble xác suất trên validation.\nChuẩn bị mã tính chỉ số theo lớp.',
 'Tạo confusion matrix, ROC và biểu diễn PCA.\nViết phần phương pháp đánh giá.',
 'Hoàn thiện luồng suy luận từ checkpoint.\nXây dựng giao diện Gradio cho tệp và ghi âm.',
 'Kiểm tra trường hợp tệp lỗi, im lặng và âm thanh dài.\nChỉnh cách hiển thị xác suất trong demo.',
 'Rà soát số liệu và các cặp nhãn nhầm lẫn.\nBiên soạn chương thực nghiệm và thảo luận.',
 'Soạn nội dung slide và kịch bản demo.\nRà soát trích dẫn, bảng và hình trong báo cáo.',
 'Tổng hợp lần chạy notebook và đối chiếu cấu hình.\nHoàn thiện báo cáo, nội dung slide và nhật ký đến 19/10/2026.'
]
start=datetime.date(2026,6,16);end=datetime.date(2026,10,19);rows=[]
for i,task in enumerate(tasks):
 a=start+datetime.timedelta(days=7*i);b=min(a+datetime.timedelta(days=6),end)
 rows.append([i+1,a.strftime('%d/%m'),b.strftime('%d/%m'),'',task,''])
table(['Tuần','Từ ngày','Đến ngày','Ngày báo cáo GVHD','Nội dung công việc','Nhận xét GVHD'],rows,[1.15,1.4,1.4,2.0,6.95,2.1],10)
p('Xác nhận của sinh viên','Content',WD_ALIGN_PARAGRAPH.RIGHT)
p('Lương Công Phú Boy                         Trần Anh Khoa','Content',WD_ALIGN_PARAGRAPH.RIGHT)

# Preserve all five source section geometries; remove page numbers in appendix.
last=d._element.body.find(qn('w:sectPr'))
last.getparent().replace(last,deepcopy(sec[-1]))
for n in list(d._element.body[-1]):
 if n.tag==qn('w:pgNumType'):d._element.body[-1].remove(n)
appendix_type=d._element.body[-1].find(qn('w:type'))
if appendix_type is not None:appendix_type.set(qn('w:val'),'nextPage')
for x in d.paragraphs:
 if x.style.name in ['Heading 1','Heading 2','Heading 3']:
  np=x._p.get_or_add_pPr().find(qn('w:numPr'))
  if np is not None:x._p.get_or_add_pPr().remove(np)
  np=OxmlElement('w:numPr');ni=OxmlElement('w:numId');ni.set(qn('w:val'),'0');np.append(ni);x._p.get_or_add_pPr().append(np)
for s in d.sections:
 for part in [s.header,s.footer]:
  for x in part.paragraphs:
   if 'Hướng dẫn và mẫu báo cáo' in x.text:replace(x,'Phân loại âm thanh đô thị sử dụng Deep Learning')
   if 'Tên sinh viên' in x.text:
    texts=list(x._p.iter(qn('w:t')))
    for j,n in enumerate(texts):
     if j<12:n.text='Lương Công Phú Boy - Trần Anh Khoa' if j==0 else ''
# Reference section uses the source lower-letter numbering and a matching PAGE field.
refs_footer=d.sections[3].footer
if not any('PAGE' in (n.text or '') for n in refs_footer._element.iter(qn('w:instrText'))):
 x=refs_footer.paragraphs[0]
 x.add_run().add_tab()
 field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),' PAGE ')
 rr=OxmlElement('w:r');tt=OxmlElement('w:t');tt.text='a';rr.append(tt);field.append(rr);x._p.append(field)
# Main header/footer are user-requested editable text; source fields and borders retained.
for rel in d.part.rels.values():
 if '/footer' in rel.reltype and not rel.is_external:
  root=rel.target_part.element
  # Last section footer reference identifies the source appendix footer.
  rid=d._element.body[-1].find(qn('w:footerReference'))
  if rid is not None and d.part.rels[rid.get(qn('r:id'))].target_part==rel.target_part:
   for field in root.findall('.//'+qn('w:fldSimple')):field.getparent().remove(field)
   # Remove entire paragraphs containing complex PAGE fields to suppress numbering.
   for x in list(root.findall(qn('w:p'))):
    if any('PAGE' in (n.text or '') for n in x.iter(qn('w:instrText'))):root.remove(x)
   if not len(root):root.append(OxmlElement('w:p'))
settings=d.settings.element
u=settings.find(qn('w:updateFields'))
if u is None:u=OxmlElement('w:updateFields');settings.append(u)
u.set(qn('w:val'),'true')
d.core_properties.title='Phân loại âm thanh đô thị sử dụng Deep Learning'
d.core_properties.author='Lương Công Phú Boy; Trần Anh Khoa'
d.core_properties.subject='Khóa luận tốt nghiệp Khoa học dữ liệu năm 2026'
work=TMP/'authored.docx';d.save(work)
# Merge only edited package parts, retaining opaque original parts byte-for-byte.
with zipfile.ZipFile(REF) as old,zipfile.ZipFile(work) as new,zipfile.ZipFile(FINAL,'w',zipfile.ZIP_DEFLATED) as out:
 for name in old.namelist():
  mutable=name in ['word/document.xml','word/_rels/document.xml.rels','word/settings.xml','docProps/core.xml','[Content_Types].xml'] or re.match(r'word/(header|footer)\d+\.xml$',name)
  out.writestr(name,new.read(name) if mutable and name in new.namelist() else old.read(name))
 for name in new.namelist():
  if name not in old.namelist():out.writestr(name,new.read(name))
with zipfile.ZipFile(REF) as a,zipfile.ZipFile(FINAL) as b:
 changed=[n for n in a.namelist() if a.read(n)!=b.read(n)]
 assert a.read('word/styles.xml')==b.read('word/styles.xml')
 assert a.read('word/numbering.xml')==b.read('word/numbering.xml')
 (TMP/'changed_parts.json').write_text(json.dumps(changed,indent=2),encoding='utf8')
print(str(FINAL));print('paragraphs',len(d.paragraphs),'tables',len(d.tables),'inline figures',len(d.inline_shapes),'sections',len(d.sections));print('changed',changed)
(TMP/'locators.json').write_text(json.dumps({'outline':outline,'figs':figs,'tables':tables},ensure_ascii=False,indent=2),encoding='utf8')
