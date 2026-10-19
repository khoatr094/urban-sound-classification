from pathlib import Path
from copy import deepcopy
import re,json,zipfile,hashlib,shutil
from lxml import etree
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH,WD_TAB_ALIGNMENT,WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from PIL import Image

ROOT=Path(__file__).resolve().parents[2];TMP=ROOT/'tmp/thesis_update'
SOURCE=ROOT/'Maubaocao_KLTN_Phan_loai_am_thanh_do_thi.docx'
SNAP=TMP/'source_snapshot.docx'
if not SNAP.exists():shutil.copy2(SOURCE,SNAP)
FINAL=ROOT/'output/Maubaocao_KLTN_Xay_dung_he_thong_phan_loai_am_thanh_do_thi.docx'
TITLE='Xây dựng hệ thống phân loại âm thanh đô thị sử dụng Deep Learning'
OLD='Phân loại âm thanh đô thị sử dụng Deep Learning'
d=Document(SNAP)
cache=json.loads((TMP/'pagecache.json').read_text(encoding='utf8')) if (TMP/'pagecache.json').exists() else json.loads((ROOT/'tmp/thesis/pagecache.json').read_text(encoding='utf8'))
loc=json.loads((ROOT/'tmp/thesis/locators.json').read_text(encoding='utf8'))
loc['outline'].insert(next(i for i,x in enumerate(loc['outline']) if x[1]=='ch5'),['4.9  Kiến trúc chi tiết của các mô hình','h4_9',1])
newfigs=[
 ('4.5','cnn4_detail.png','Các tầng và kích thước đặc trưng của CNN4','Hình 4.5 mô tả CNN4 với đầu vào một kênh gồm 64 dải Mel và 251 khung thời gian. Mỗi khối có hai lớp tích chập trước khi giảm kích thước bằng max pooling. Bốn khối tạo bản đồ 256 × 4 × 15; pooling trung bình toàn cục đưa bản đồ này thành vector 256 chiều. Các kích thước được xác định theo tiền xử lý và kiến trúc của notebook đã chạy.'),
 ('4.6','rescnn_detail.png','Kiến trúc ResCNN và nhánh pooling kép','Ở Hình 4.6, stride 2 trong lớp tích chập đầu của mỗi khối làm giảm độ phân giải của đặc trưng. Với padding 1, chiều lẻ được làm tròn lên khi giảm kích thước, vì vậy chiều thời gian lần lượt là 251, 126, 63, 32 và 16. Nhánh pooling trung bình giữ mức đáp ứng chung, còn nhánh pooling cực đại giữ đáp ứng mạnh nhất. Hai vector được ghép trước tầng phân loại.'),
 ('4.7','resblock_detail.png','Chi tiết khối phần dư và cơ chế chú ý kênh SE','Hình 4.7 tách riêng nhánh biến đổi và nhánh tắt. Khối SE tạo một trọng số cho mỗi kênh rồi nhân trọng số với đặc trưng của nhánh chính. Phép cộng chỉ thực hiện khi hai nhánh có cùng kích thước; nhánh tắt dùng phép chiếu 1 × 1 nếu cần. Cấu trúc này vận dụng học phần dư và chú ý kênh đã công bố [2], [3], không phải một cơ chế mới của đề tài.'),
 ('4.8','ast_detail.png','Kiến trúc AST và một tầng Transformer encoder','Hình 4.8 được xây dựng theo cấu hình AST của dự án: patch 16 × 16, stride 10, chiều ẩn 768, 12 tầng encoder và 12 đầu attention. Đầu vào 1.024 × 128 tạo 101 vị trí thời gian và 12 vị trí tần số, tương ứng 1.212 patch. Sau khi thêm CLS và DIST, chuỗi có 1.214 token. Hai token đặc biệt được tổng hợp trước đầu phân loại bảy lớp [1], [4].')
]
ix=next(i for i,x in enumerate(loc['figs']) if x[1]=='fig5_1')
for j,(num,_,caption,_) in enumerate(newfigs):loc['figs'].insert(ix+j,[f'Hình {num} {caption}','fig'+num.replace('.','_')])
lookup={re.sub(r'\s+','',title).casefold():key for title,key,*_ in loc['outline']+loc['figs']+loc['tables']}
lookup.update({re.sub(r'\s+','',a).casefold():b for a,b in [('TÀI LIỆU THAM KHẢO','refs'),('PHỤ LỤC','appendix'),('NHẬT KÝ LÀM VIỆC','journal'),('Bảng A.1 Các tệp dùng để kiểm tra số liệu','tabA1'),('Bảng B.1 Nhật ký công việc theo tuần','tabB1')]})

def reset(p,text):
 rp=deepcopy(p.runs[0]._r.rPr) if p.runs and p.runs[0]._r.rPr is not None else None
 for x in list(p._p):
  if x.tag!=qn('w:pPr'):p._p.remove(x)
 r=p.add_run(text)
 if rp is not None:r._r.insert(0,rp)
def replace_across(part,old,new):
 for paragraph in part.iter(qn('w:p')):
  nodes=list(paragraph.iter(qn('w:t')));full=''.join(x.text or '' for x in nodes)
  for match in reversed(list(re.finditer(re.escape(old),full,re.I))):
   pos=0;first=last=None;startoff=endoff=0
   for i,x in enumerate(nodes):
    ln=len(x.text or '')
    if first is None and pos<=match.start()<pos+ln:first=i;startoff=match.start()-pos
    if pos<match.end()<=pos+ln:last=i;endoff=match.end()-pos;break
    pos+=ln
   if first is None or last is None:continue
   before=(nodes[first].text or '')[:startoff];after=(nodes[last].text or '')[endoff:]
   nodes[first].text=before+new+after
   for i in range(first+1,last+1):nodes[i].text=''

for pi,p in enumerate(d.paragraphs):
 if p.text.replace('\n',' ').strip()==OLD.upper():reset(p,'XÂY DỰNG HỆ THỐNG PHÂN LOẠI\nÂM THANH ĐÔ THỊ\nSỬ DỤNG DEEP LEARNING')
 elif p.text.replace('\n',' ').strip()=='URBAN SOUND CLASSIFICATION USING DEEP LEARNING':
  reset(p,'DEVELOPMENT OF AN URBAN SOUND\nCLASSIFICATION SYSTEM\nUSING DEEP LEARNING' if pi<72 else 'DEVELOPMENT OF AN URBAN SOUND\nCLASSIFICATION SYSTEM USING DEEP LEARNING')
 elif p.text.startswith('These findings are limited to a seven-class'):
  reset(p,'The study uses one seven-class test fold and is not a ten-fold UrbanSound8K benchmark. Further work should cover multiple folds and seeds, out-of-class evaluation and a version-locked deployment package.')
replace_across(d._element,OLD,TITLE)
# Keep the longer title on each original cover by reducing only blank spacer height.
for p in d.paragraphs[:72]:
 if 'XÂY DỰNG HỆ THỐNG PHÂN LOẠI' not in p.text and 'DEVELOPMENT OF AN URBAN SOUND' not in p.text:continue
 ps=d.paragraphs;idx=next(i for i,x in enumerate(ps) if x._p is p._p)
 stop=next((i for i in range(idx+1,len(ps)) if 'THÁNH' in ps[i].text or 'THÀNH PHỐ HỒ CHÍ MINH, THÁNG' in ps[i].text or 'HO CHI MINH CITY – OCTOBER' in ps[i].text),idx+12)
 blanks=[x for x in ps[idx+1:stop] if not x.text.strip() and not x._p.xpath('.//w:drawing|.//w:pict|.//w:sectPr|.//w:br')]
 for x in blanks[1::2][:3]:
  x.paragraph_format.space_before=Pt(0);x.paragraph_format.space_after=Pt(0);x.paragraph_format.line_spacing=Pt(2)
  for rr in x.runs:rr.font.size=Pt(2)
for rel in d.part.rels.values():
 if not rel.is_external and ('/header' in rel.reltype or '/footer' in rel.reltype):replace_across(rel.target_part.element,OLD,TITLE)

counter=max([int(x.get(qn('w:id'))) for x in d._element.iter(qn('w:bookmarkStart'))]+[1000])+1
def bookmark(p,name):
 global counter
 a=OxmlElement('w:bookmarkStart');a.set(qn('w:id'),str(counter));a.set(qn('w:name'),name)
 b=OxmlElement('w:bookmarkEnd');b.set(qn('w:id'),str(counter));counter+=1
 p._p.insert(1,a);p._p.append(b)
def insertpara(before,text='',style='Content'):
 p=d.add_paragraph(text,style);before.addprevious(p._p)
 p.paragraph_format.widow_control=True
 return p
def stopnum(p):
 pp=p._p.get_or_add_pPr();old=pp.find(qn('w:numPr'))
 if old is not None:pp.remove(old)
 n=OxmlElement('w:numPr');i=OxmlElement('w:numId');i.set(qn('w:val'),'0');n.append(i);pp.append(n)

# Native Office Math objects: fractions, limits, subscripts and powers remain editable.
def mr(text):
 r=OxmlElement('m:r');wp=OxmlElement('w:rPr');fn=OxmlElement('w:rFonts');fn.set(qn('w:ascii'),'Cambria Math');fn.set(qn('w:hAnsi'),'Cambria Math');wp.append(fn)
 size=OxmlElement('w:sz');size.set(qn('w:val'),'26');wp.append(size);r.append(wp)
 t=OxmlElement('m:t');t.text=str(text);t.set(qn('xml:space'),'preserve');r.append(t);return r
def seq(*items):return [mr(x) if isinstance(x,(str,int)) else x for x in items]
def container(tag,items):
 e=OxmlElement(tag)
 for x in items if isinstance(items,list) else [items]:e.append(mr(x) if isinstance(x,(str,int)) else x)
 return e
def frac(a,b):
 e=OxmlElement('m:f');pr=OxmlElement('m:fPr');typ=OxmlElement('m:type');typ.set(qn('m:val'),'bar');pr.append(typ);e.append(pr);e.append(container('m:num',a));e.append(container('m:den',b));return e
def sub(a,b):
 e=OxmlElement('m:sSub');e.append(container('m:e',a));e.append(container('m:sub',b));return e
def sup(a,b):
 e=OxmlElement('m:sSup');e.append(container('m:e',a));e.append(container('m:sup',b));return e
def sqrt(a):
 e=OxmlElement('m:rad');pr=OxmlElement('m:radPr');v=OxmlElement('m:degHide');v.set(qn('m:val'),'1');pr.append(v);e.append(pr);e.append(OxmlElement('m:deg'));e.append(container('m:e',a));return e
def summ(lower,upper,body):
 e=OxmlElement('m:nary');pr=OxmlElement('m:naryPr');ch=OxmlElement('m:chr');ch.set(qn('m:val'),'∑');pr.append(ch);lm=OxmlElement('m:limLoc');lm.set(qn('m:val'),'undOvr');pr.append(lm);e.append(pr);e.append(container('m:sub',lower));e.append(container('m:sup',upper));e.append(container('m:e',body));return e
def mathline(before,parts,label):
 p=insertpara(before,'','Normal');pf=p.paragraph_format
 pf.first_line_indent=Cm(0);pf.left_indent=Cm(0);pf.right_indent=Cm(0);pf.space_before=Pt(10);pf.space_after=Pt(10);pf.line_spacing=1.3
 pf.tab_stops.add_tab_stop(Cm(7.0),WD_TAB_ALIGNMENT.CENTER);pf.tab_stops.add_tab_stop(Cm(15.0),WD_TAB_ALIGNMENT.RIGHT)
 p.add_run('\t');m=OxmlElement('m:oMath')
 for item in parts:m.append(mr(item) if isinstance(item,(str,int)) else item)
 p._p.append(m);p.add_run('\t('+label+')')
 bookmark(p,'eq'+label.replace('.','_'))
 return p
def defs(before,heading,items):
 x=insertpara(before,heading);x.paragraph_format.keep_with_next=True
 for symbol,text in items:
  x=insertpara(before,'');pos=0
  for match in re.finditer(r'(μ|σ|TP|FP|FN|F1|P|R|N|p|w)_(train|đúng|ens|c|m)',symbol):
   r=x.add_run(symbol[pos:match.start()]+match.group(1));r.bold=True
   r=x.add_run(match.group(2));r.bold=True;r.font.subscript=True;pos=match.end()
  r=x.add_run(symbol[pos:]+': ');r.bold=True;x.add_run(text)
  x.paragraph_format.left_indent=Cm(.45);x.paragraph_format.first_line_indent=Cm(-.45)
  x.paragraph_format.space_before=Pt(3);x.paragraph_format.space_after=Pt(0)

specs={
 '2.1':([
  (lambda:seq('X(m,k) = ',summ('n = 0','N − 1',seq('x[n + mH]w[n]',sup('e',seq('−j',frac('2πkn','N')))))),'2.1')
 ],'Các thành phần của phép biến đổi Fourier ngắn hạn gồm:',[
 ('x[n + mH]','mẫu tín hiệu tại vị trí n trong khung thứ m. Chỉ số m xác định khung thời gian, còn n chạy từ 0 đến N − 1 trong khung.'),
 ('w[n]','giá trị hàm cửa sổ tại mẫu n; phép nhân với cửa sổ làm giảm gián đoạn ở hai đầu đoạn tín hiệu.'),
 ('H và N','H là số mẫu dịch giữa hai khung liên tiếp; N là số điểm biến đổi. Trong nhánh CNN, H = 256 và N = 1.024.'),
 ('k và j','k là chỉ số bin tần số; j là đơn vị ảo với j² = −1. Thương số trong số mũ xác định pha của hàm cơ sở Fourier.'),
 ('X(m,k)','hệ số phức tại khung m và bin k. Độ lớn mô tả biên độ của thành phần tần số; bình phương độ lớn tạo phổ công suất.')]),
 '2.2':([
  (lambda:seq('M(b,m) = ',summ('k = 0','K − 1',seq(sub('H','b'),'(k)',sup('|X(m,k)|','2')))),'2.2a'),
  (lambda:seq('L(b,m) = 10 ',sub('log','10'),'(max(M(b,m), ε))'),'2.2b')
 ],'Hai bước trong phép tạo phổ Log-Mel được hiểu như sau:',[
 ('M(b,m)','năng lượng của dải Mel thứ b tại khung m, được tính bằng tổng năng lượng các bin tần số có trọng số.'),
 ('Hᵦ(k)','trọng số của bộ lọc Mel thứ b tại bin k. Các bộ lọc chồng lấn và gom nhiều bin Fourier vào một dải Mel.'),
 ('K và |X(m,k)|²','K là số bin tần số được đưa vào bộ lọc; |X(m,k)|² là công suất tại bin k và khung m.'),
 ('L(b,m), log₁₀ và ε','L là giá trị sau phép log; hệ số 10 áp dụng cho đại lượng công suất. ε là số dương rất nhỏ để tránh lấy log của 0. Mã CNN còn giới hạn dải động bằng top_db = 80.')]),
 '2.3':([
  (lambda:seq(sub('h','out'),' = φ(F(',sub('h','in'),') + P(',sub('h','in'),'))'),'2.3')
 ],'Ý nghĩa của từng nhánh trong công thức phần dư:',[
 ('hᵢₙ và hₒᵤₜ','lần lượt là tensor đầu vào và đầu ra của khối. Tensor gồm chiều kênh, tần số và thời gian.'),
 ('F(hᵢₙ)','nhánh biến đổi học được từ các lớp tích chập, chuẩn hóa, chú ý kênh và dropout trong khối ResCNN.'),
 ('P(hᵢₙ)','nhánh tắt. P là ánh xạ đồng nhất nếu kích thước giữ nguyên, hoặc phép chiếu 1 × 1 kèm chuẩn hóa nếu cần đổi kích thước.'),
 ('Dấu cộng và φ','phép cộng thực hiện theo từng phần tử của hai tensor cùng kích thước. φ là hàm kích hoạt sau phép cộng; mô hình của đề tài dùng GELU.')]),
 '2.4':([
  (lambda:seq('Attention(Q,K,V) = softmax(',frac(seq('Q',sup('K','T')),sqrt(sub('d','k'))),')V'),'2.4')
 ],'Các thành phần của một đầu self-attention:',[
 ('Q, K và V','các ma trận query, key và value, được tạo bằng ba phép chiếu tuyến tính học được từ chuỗi token đầu vào.'),
 ('Kᵀ và QKᵀ','T biểu thị phép chuyển vị. Tích QKᵀ cho điểm tương đồng giữa query tại một token và key tại các token còn lại.'),
 ('dₖ và căn bậc hai của dₖ','dₖ là số chiều key của một đầu attention. Chia điểm tương đồng cho căn bậc hai của dₖ giúp kiểm soát thang giá trị trước softmax.'),
 ('softmax và phép nhân V','softmax được áp dụng trên các key để tạo trọng số có tổng bằng 1. Nhân với V tạo tổng có trọng số của các vector value. AST của dự án có 12 đầu; mỗi đầu có 64 chiều khi chiều ẩn bằng 768.')]),
 '3.1':([
  (lambda:seq('Z(b,t) = ',frac(seq('L(b,t) − ',sub('μ','train'),'(b)'),seq('max(',sub('σ','train'),'(b), ε)'))),'3.1')
 ],'Phân tích phép chuẩn hóa theo từng dải Mel:',[
 ('b, t và L(b,t)','b xác định dải Mel, t xác định khung thời gian, còn L là giá trị Log-Mel trước chuẩn hóa.'),
 ('μ_train(b)','trung bình của dải b, tính trên các mẫu và các khung thời gian thuộc train. Phần tử này được trừ ở tử số để đưa dải về gần trung bình 0.'),
 ('σ_train(b)','độ lệch chuẩn của cùng dải trên train. Mẫu số đưa các dải về thang biến thiên tương đối đồng đều.'),
 ('max và ε','mẫu số lấy giá trị lớn hơn giữa độ lệch chuẩn và ε để tránh chia cho 0. Validation và test dùng lại thống kê train; Z là đặc trưng được đưa vào CNN.')]),
 '3.2':([
  (lambda:seq('Accuracy = ',frac(sub('N','đúng'),'N')),'3.2a'),
  (lambda:seq(sub('P','c'),' = ',frac(sub('TP','c'),seq(sub('TP','c'),' + ',sub('FP','c')))),'3.2b'),
  (lambda:seq(sub('R','c'),' = ',frac(sub('TP','c'),seq(sub('TP','c'),' + ',sub('FN','c')))),'3.2c')
 ],'Các đại lượng đếm trong ba công thức đánh giá:',[
 ('N_đúng và N','N_đúng là số mẫu có nhãn dự đoán trùng nhãn thật; N là tổng số mẫu đánh giá. Accuracy lấy tỉ lệ giữa hai số này.'),
 ('TP_c','số mẫu thực sự thuộc lớp c và được dự đoán đúng là lớp c.'),
 ('FP_c','số mẫu thuộc lớp khác nhưng bị dự đoán là c. Mẫu số của precision là toàn bộ số mẫu được dự đoán là c.'),
 ('FN_c','số mẫu thực sự thuộc lớp c nhưng bị gán sang lớp khác. Mẫu số của recall là toàn bộ số mẫu thật thuộc c.'),
 ('P_c và R_c','lần lượt là precision và recall của lớp c. Khi một mẫu số bằng 0, cần dùng quy ước zero_division thống nhất; quá trình tính chỉ số của dự án đặt giá trị tương ứng bằng 0.')]),
 '3.3':([
  (lambda:seq(sub('F1','c'),' = ',frac(seq('2',sub('P','c'),sub('R','c')),seq(sub('P','c'),' + ',sub('R','c')))),'3.3a'),
  (lambda:seq('Macro-F1 = ',frac('1','C'),summ('c = 1','C',sub('F1','c'))),'3.3b')
 ],'Cách kết hợp precision và recall thành F1:',[
 ('2P_cR_c','tử số của F1, bằng hai lần tích precision và recall của lớp c. Nếu một trong hai bằng 0, F1 cũng bằng 0.'),
 ('P_c + R_c','mẫu số của trung bình điều hòa. F1 chỉ cao khi cả precision và recall cùng cao; trường hợp cả hai bằng 0 được quy ước F1 bằng 0.'),
 ('C và phép lấy tổng','C là số lớp, bằng 7 trong đề tài. Tổng F1 của C lớp được nhân với phân số một phần C để mỗi lớp đóng góp ngang nhau.'),
 ('Đơn vị trình bày','các tỉ lệ có giá trị từ 0 đến 1 trước khi đổi sang phần trăm. Macro-F1 không bằng F1 tính từ macro-precision và macro-recall trong mọi trường hợp.')]),
 '4.1':([
  (lambda:seq(sub('p','ens'),'(c|x) = ',summ('m = 1','M',seq(sub('w','m'),sub('p','m'),'(c|x)'))),'4.1a'),
  (lambda:seq(sub('w','m'),' ≥ 0,    ',summ('m = 1','M',sub('w','m')),' = 1'),'4.1b')
 ],'Ý nghĩa các thành phần của ensemble xác suất:',[
 ('x và c','x là cùng một mẫu âm thanh cần phân loại; c là một trong bảy lớp mục tiêu. Thứ tự lớp phải khớp giữa các mô hình.'),
 ('p_m(c|x)','xác suất softmax mà mô hình thứ m gán cho lớp c trên mẫu x.'),
 ('w_m và M','w_m là trọng số của mô hình m; M là số mô hình kết hợp, bằng 3 trong thực nghiệm. Trọng số không âm và có tổng bằng 1.'),
 ('p_ens(c|x)','tổng có trọng số của các xác suất thành phần. Nếu các vector thành phần có tổng bằng 1, điều kiện trọng số giúp vector kết hợp cũng có tổng bằng 1. Nhãn cuối là lớp có xác suất kết hợp lớn nhất.')])
}
done=[]
for p in list(d.paragraphs):
 m=re.search(r'\((2\.[1-4]|3\.[1-3]|4\.1)\)\s*$',p.text)
 if not m or not any(x in p.text for x in [' = ','=']):continue
 key=m.group(1)
 if key in done:continue
 rows,heading,items=specs[key]
 for fn,label in rows:mathline(p._p,fn(),label)
 defs(p._p,heading,items)
 p._p.getparent().remove(p._p);done.append(key)
assert set(done)==set(specs),(done,list(specs))

# Supplement the existing overview figures with layer-level architectural diagrams.
before=next(p._p for p in d.paragraphs if p.style.name=='Heading 1' and 'CHƯƠNG 5' in p.text)
h=insertpara(before,'4.9  Kiến trúc chi tiết của các mô hình','Heading 2');stopnum(h);bookmark(h,'h4_9');h.paragraph_format.keep_with_next=True
insertpara(before,'Các sơ đồ dưới đây làm rõ thành phần bên trong từng mô hình, bổ sung cho luồng tổng quát ở các mục trước. Kích thước tensor của hai CNN được xác định với đặc trưng 64 × 251. Kích thước và số tầng của AST lấy từ cấu hình checkpoint dùng trong dự án; các hình là sơ đồ kiến trúc, không phải kết quả đo hiệu năng.')
for num,path,title,text in newfigs:
 im=Image.open(TMP/path);w,ht=im.size;scale=min(15.5/w,15.3/ht)
 p=insertpara(before,'','Normal');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
 pf=p.paragraph_format;pf.first_line_indent=Cm(0);pf.left_indent=Cm(0);pf.right_indent=Cm(0);pf.line_spacing=1;pf.space_before=Pt(8);pf.space_after=Pt(0);pf.keep_with_next=True
 p.add_run().add_picture(str(TMP/path),width=Cm(w*scale),height=Cm(ht*scale))
 p=insertpara(before,f'Hình {num} {title}','Caption');p.alignment=WD_ALIGN_PARAGRAPH.CENTER;bookmark(p,'fig'+num.replace('.','_'))
 p.paragraph_format.first_line_indent=Cm(0);p.paragraph_format.keep_with_next=True;p.paragraph_format.space_after=Pt(6)
 for r in p.runs:r.font.size=Pt(13);r.italic=False
 p=insertpara(before,'Nguồn: nhóm xây dựng theo kiến trúc và cấu hình của notebook.','Normal');p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Cm(0);p.paragraph_format.line_spacing=1
 for r in p.runs:r.font.size=Pt(10);r.italic=True
 insertpara(before,text)

# Identify the front-matter lists before turning their entries into hyperlinks.
allp=d.paragraphs
start=next(i for i,p in enumerate(allp) if p.text.strip()=='MỤC LỤC')
end=next(i for i,p in enumerate(allp) if p.style.name=='Heading 1' and 'CHƯƠNG 1' in p.text)
entry=next(p for p in allp[start:end] if p.text.startswith('CHƯƠNG 5'))
insertpara(entry._p,'4.9  Kiến trúc chi tiết của các mô hình\t'+str(cache.get('h4_9','…')),'Normal')
entry=next(p for p in allp[start:end] if p.text.startswith('Hình 5.1 '))
for num,_,title,_ in newfigs:insertpara(entry._p,f'Hình {num} {title}\t'+str(cache.get('fig'+num.replace('.','_'),'…')),'Normal')

def wr(text,props=None):
 r=OxmlElement('w:r')
 if props is not None:r.append(deepcopy(props))
 t=OxmlElement('w:t');t.text=text;t.set(qn('xml:space'),'preserve');r.append(t);return r
def nav(p,title,key,level):
 # Entire entry links to its bookmark; verified page numbers survive Word/WPS rendering.
 reset(p,'');pf=p.paragraph_format;pf.first_line_indent=Cm(0);pf.left_indent=Cm(.45*level);pf.line_spacing=1.15;pf.space_before=Pt(0);pf.space_after=Pt(3)
 pf.tab_stops.add_tab_stop(Cm(15-.45*level),WD_TAB_ALIGNMENT.RIGHT,WD_TAB_LEADER.DOTS)
 link=OxmlElement('w:hyperlink');link.set(qn('w:anchor'),key);link.set(qn('w:history'),'1');link.set(qn('w:tooltip'),'Chuyển đến '+title)
 link.append(wr(title));r=OxmlElement('w:r');r.append(OxmlElement('w:tab'));link.append(r)
 value=str(cache.get(key,'…'))
 if value:
  link.append(wr(value))
 p._p.append(link)

inlists=False;navcount=0
for p in d.paragraphs:
 if p.text.strip()=='MỤC LỤC':inlists=True;continue
 if p.style.name=='Heading 1' and 'CHƯƠNG 1' in p.text:inlists=False
 if not inlists or '\t' not in p.text:continue
 title=p.text.split('\t')[0].strip();key=lookup.get(re.sub(r'\s+','',title).casefold())
 if not key:continue
 level=1 if re.match(r'^\d+\.\d+\s',title) else 0
 nav(p,title,key,level);navcount+=1

# Bibliography entries are full-entry external hyperlinks to verified source URLs.
refcount=0
for p in d.paragraphs:
 if not re.match(r'^\[\d+\]',p.text):continue
 urls=re.findall(r'https?://[^\s]+',p.text)
 if not urls:continue
 url=urls[0].rstrip('.,);');n=int(re.match(r'^\[(\d+)\]',p.text).group(1));bookmark(p,f'ref{n}')
 rel=d.part.relate_to(url,RT.HYPERLINK,is_external=True)
 link=OxmlElement('w:hyperlink');link.set(qn('r:id'),rel);link.set(qn('w:history'),'1');link.set(qn('w:tooltip'),url)
 for x in list(p._p):
  if x.tag==qn('w:r'):p._p.remove(x);link.append(x)
 p._p.append(link);refcount+=1

d.core_properties.title=TITLE
settings=d.settings.element;u=settings.find(qn('w:updateFields'))
if u is None:u=OxmlElement('w:updateFields');settings.append(u)
u.set(qn('w:val'),'true')
work=TMP/'authored.docx';d.save(work)
with zipfile.ZipFile(SNAP) as old,zipfile.ZipFile(work) as new,zipfile.ZipFile(FINAL,'w',zipfile.ZIP_DEFLATED) as out:
 for name in old.namelist():
  mutable=name in ['word/document.xml','word/_rels/document.xml.rels','word/settings.xml','docProps/core.xml','[Content_Types].xml'] or re.match(r'word/(header|footer)\d+\.xml$',name)
  out.writestr(name,new.read(name) if mutable and name in new.namelist() else old.read(name))
 for name in new.namelist():
  if name not in old.namelist():out.writestr(name,new.read(name))
(TMP/'locators.json').write_text(json.dumps(loc,ensure_ascii=False,indent=2),encoding='utf8')
(TMP/'manifest.json').write_text(json.dumps({'source':str(SOURCE),'source_sha256':hashlib.sha256(SNAP.read_bytes()).hexdigest(),'output':str(FINAL),'equation_groups':done,'internal_links':navcount,'web_links':refcount,'new_architecture_figures':4},ensure_ascii=False,indent=2),encoding='utf8')
print(FINAL);print('equation groups',done,'internal links',navcount,'web links',refcount,'total inline images',len(d.inline_shapes))
